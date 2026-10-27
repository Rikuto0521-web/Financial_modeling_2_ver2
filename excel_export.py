# excel_export.py
import io
 
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
 
from model import NUM_YEARS, YEARS, _tile
 
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")  # 入力値
BLACK = Font(name=FONT)  # 数式
GREEN = Font(name=FONT, color="008000")  # 他シートへのリンク
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", start_color="2B5C8F")
INPUT_FILL = PatternFill("solid", start_color="FFFF99")
 
NUM = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
 
A = "'前提'"  # 前提シートの参照名
 
 
def _header(ws, row, labels, first_col=2):
    ws.cell(row=row, column=1, value="").fill = HEAD_FILL
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=first_col + i, value=lab)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
 
 
def _label(ws, row, text, bold=False):
    c = ws.cell(row=row, column=1, value=text)
    c.font = BOLD if bold else BLACK
 
 
def _put(ws, ref, value, font=BLACK, fmt=NUM, fill=None):
    c = ws[ref]
    c.value = value
    c.font = font
    c.number_format = fmt
    if fill:
        c.fill = fill
 
 
def build_workbook(p, extra_sheets=None):
    """
    前提シートの数字を変えると、全シートが再計算される数式入りのExcelを作る。
    p: model.simulate() に渡すものと同じ辞書
    extra_sheets: {"シート名": DataFrame} 値のまま貼り付けるシート
    """
    wb = Workbook()
 
    # =====================================================
    # 前提シート
    # =====================================================
    ws = wb.active
    ws.title = "前提"
    ws["A1"] = "前提条件 (青字・黄色セルが入力値。変更すると全シートが再計算されます)"
    ws["A1"].font = TITLE
 
    scalars = [
        (3, "実効税率", p["tax_rate"], PCT),
        (4, "設備投資率 (対売上)", p["capex_pct"], PCT),
        (5, "減価償却費率 (対売上)", p["da_pct"], PCT),
        (6, "売上債権回転日数 (日)", p["ar_days"], "0"),
        (7, "棚卸資産回転日数 (日)", p["inv_days"], "0"),
        (8, "買入債務回転日数 (日)", p["ap_days"], "0"),
        (9, "借入金利息率", p["interest_rate"], PCT),
        (10, "年間返済率 (対 初期有利子負債)", p["repay_rate"], PCT),
        (11, "配当性向 (対 当期純利益)", p["payout"], PCT),
        (12, "WACC", p["wacc"], '0.00%'),
        (13, "永続成長率", p["perp_growth"], '0.00%'),
        (14, "エグジットEV/EBITDA倍率", p["exit_multiple"], MULT),
        (15, "発行済株式数 (万株)", p["shares_outstanding"], NUM),
        (16, "基準売上高 (Yr0)", p["base_rev"], NUM),
        (17, "初期現金預金 (Yr0)", p["base_cash"], NUM),
        (18, "初期有形固定資産 (Yr0)", p["base_ppe"], NUM),
        (19, "初期有利子負債 (Yr0)", p["base_debt"], NUM),
    ]
    for row, label, val, fmt in scalars:
        _label(ws, row, label)
        _put(ws, f"B{row}", val, BLUE, fmt, INPUT_FILL)
 
    ws["A21"] = "年別ドライバ"
    ws["A21"].font = BOLD
    _header(ws, 22, YEARS)
    g = _tile(p["rev_growth"])
    c = _tile(p["cogs_pct"])
    s = _tile(p["sga_pct"])
    _label(ws, 23, "売上高成長率")
    _label(ws, 24, "売上原価率")
    _label(ws, 25, "販管費率")
    _put(ws, "B24", p["base_cogs_pct"], BLUE, PCT, INPUT_FILL)
    _put(ws, "B25", p["base_sga_pct"], BLUE, PCT, INPUT_FILL)
    for i in range(NUM_YEARS):
        col = get_column_letter(3 + i)
        _put(ws, f"{col}23", g[i], BLUE, PCT, INPUT_FILL)
        _put(ws, f"{col}24", c[i], BLUE, PCT, INPUT_FILL)
        _put(ws, f"{col}25", s[i], BLUE, PCT, INPUT_FILL)
 
    ws["A27"] = "注: 金額の単位は百万円。Yr0の株主資本は、貸借が一致するよう資産−負債で自動計算しています。"
    ws["A27"].font = Font(name=FONT, italic=True)
    ws.column_dimensions["A"].width = 38
    for i in range(2, 8):
        ws.column_dimensions[get_column_letter(i)].width = 14
 
    # =====================================================
    # 損益計算書  (列 B..G = Yr0..Yr5)
    # =====================================================
    wi = wb.create_sheet("損益計算書")
    wi["A1"] = "損益計算書 (百万円)"
    wi["A1"].font = TITLE
    _header(wi, 2, YEARS)
    is_rows = {
        3: "売上高", 4: "売上原価", 5: "売上総利益", 6: "販管費",
        7: "営業利益 (EBIT)", 8: "支払利息", 9: "税引前利益 (EBT)",
        10: "期首 繰越欠損金", 11: "欠損金使用額", 12: "課税所得",
        13: "法人税等", 14: "期末 繰越欠損金", 15: "当期純利益",
    }
    for r, lab in is_rows.items():
        _label(wi, r, lab, bold=r in (5, 7, 15))
 
    # =====================================================
    # 貸借対照表
    # =====================================================
    wbs = wb.create_sheet("貸借対照表")
    wbs["A1"] = "貸借対照表 (百万円)"
    wbs["A1"].font = TITLE
    _header(wbs, 2, YEARS)
    bs_rows = {
        3: "現金預金", 4: "売上債権", 5: "棚卸資産", 6: "有形固定資産 (純額)",
        7: "資産合計", 8: "買入債務", 9: "有利子負債", 10: "負債合計",
        11: "株主資本 (利益余剰金含む)", 12: "負債・純資産合計",
        13: "BSバランスチェック (差額)",
    }
    for r, lab in bs_rows.items():
        _label(wbs, r, lab, bold=r in (7, 10, 12))
 
    # =====================================================
    # キャッシュフロー
    # =====================================================
    wc = wb.create_sheet("キャッシュフロー")
    wc["A1"] = "キャッシュフロー計算書 (百万円)"
    wc["A1"].font = TITLE
    _header(wc, 2, YEARS)
    cf_rows = {
        3: "当期純利益", 4: "減価償却費", 5: "運転資本の増減 (増加は−)",
        6: "営業CF", 7: "設備投資 (CapEx)", 8: "投資CF",
        9: "借入金返済", 10: "配当支払", 11: "財務CF",
        12: "期末現金増減", 13: "NOPAT", 14: "フリーキャッシュフロー (FCF)",
    }
    for r, lab in cf_rows.items():
        _label(wc, r, lab, bold=r in (6, 8, 11, 14))
 
    I, B, C = "'損益計算書'", "'貸借対照表'", "'キャッシュフロー'"
 
    # ---- Yr0 (列B) ----
    _put(wi, "B3", f"={A}!B16", GREEN)
    _put(wi, "B4", f"=B3*{A}!B24")
    _put(wi, "B5", "=B3-B4")
    _put(wi, "B6", f"=B3*{A}!B25")
    _put(wi, "B7", "=B5-B6")
    _put(wi, "B8", f"={B}!B9*{A}!$B$9")
    _put(wi, "B9", "=B7-B8")
    _put(wi, "B10", 0)
    _put(wi, "B11", "=MIN(B10,MAX(B9,0))")
    _put(wi, "B12", "=MAX(B9,0)-B11")
    _put(wi, "B13", f"=B12*{A}!$B$3")
    _put(wi, "B14", "=B10-B11+MAX(-B9,0)")
    _put(wi, "B15", "=B9-B13")
 
    _put(wbs, "B3", f"={A}!B17", GREEN)
    _put(wbs, "B6", f"={A}!B18", GREEN)
    _put(wbs, "B9", f"={A}!B19", GREEN)
    _put(wbs, "B11", "=B7-B10")
 
    # ---- 全年共通の式 (B..G) ----
    for i in range(NUM_YEARS + 1):
        col = get_column_letter(2 + i)
        prev = get_column_letter(1 + i)
        # 貸借対照表: 運転資本項目・合計
        _put(wbs, f"{col}4", f"={I}!{col}3*{A}!$B$6/365")
        _put(wbs, f"{col}5", f"={I}!{col}4*{A}!$B$7/365")
        _put(wbs, f"{col}7", f"=SUM({col}3:{col}6)")
        _put(wbs, f"{col}8", f"={I}!{col}4*{A}!$B$8/365")
        _put(wbs, f"{col}10", f"={col}8+{col}9")
        _put(wbs, f"{col}12", f"={col}10+{col}11")
        _put(wbs, f"{col}13", f"={col}7-{col}12")
 
        if i == 0:
            continue
        # 前提シートの年別ドライバは列C..G が Yr1..Yr5 (Yr0はB)
        # 損益計算書
        _put(wi, f"{col}3", f"={prev}3*(1+{A}!{col}23)")
        _put(wi, f"{col}4", f"={col}3*{A}!{col}24")
        _put(wi, f"{col}5", f"={col}3-{col}4")
        _put(wi, f"{col}6", f"={col}3*{A}!{col}25")
        _put(wi, f"{col}7", f"={col}5-{col}6")
        _put(wi, f"{col}8", f"={B}!{prev}9*{A}!$B$9")
        _put(wi, f"{col}9", f"={col}7-{col}8")
        _put(wi, f"{col}10", f"={prev}14")
        _put(wi, f"{col}11", f"=MIN({col}10,MAX({col}9,0))")
        _put(wi, f"{col}12", f"=MAX({col}9,0)-{col}11")
        _put(wi, f"{col}13", f"={col}12*{A}!$B$3")
        _put(wi, f"{col}14", f"={col}10-{col}11+MAX(-{col}9,0)")
        _put(wi, f"{col}15", f"={col}9-{col}13")
        # キャッシュフロー
        _put(wc, f"{col}3", f"={I}!{col}15")
        _put(wc, f"{col}4", f"={I}!{col}3*{A}!$B$5")
        _put(
            wc, f"{col}5",
            f"=-(({B}!{col}4-{B}!{prev}4)+({B}!{col}5-{B}!{prev}5)"
            f"-({B}!{col}8-{B}!{prev}8))",
        )
        _put(wc, f"{col}6", f"=SUM({col}3:{col}5)")
        _put(wc, f"{col}7", f"=-{I}!{col}3*{A}!$B$4")
        _put(wc, f"{col}8", f"={col}7")
        _put(wc, f"{col}9", f"=-MIN({B}!{prev}9,{A}!$B$19*{A}!$B$10)")
        _put(wc, f"{col}10", f"=-MAX({I}!{col}15,0)*{A}!$B$11")
        _put(wc, f"{col}11", f"={col}9+{col}10")
        _put(wc, f"{col}12", f"={col}6+{col}8+{col}11")
        _put(wc, f"{col}13", f"={I}!{col}7-MAX({I}!{col}7,0)*{A}!$B$3")
        _put(wc, f"{col}14", f"={col}13+{col}4+{col}7+{col}5")
        # 貸借対照表: 前年からの積み上げ項目
        _put(wbs, f"{col}3", f"={prev}3+{C}!{col}12")
        _put(
            wbs, f"{col}6",
            f"={prev}6+{I}!{col}3*{A}!$B$4-{I}!{col}3*{A}!$B$5",
        )
        _put(wbs, f"{col}9", f"={prev}9+{C}!{col}9")
        _put(wbs, f"{col}11", f"={prev}11+{I}!{col}15+{C}!{col}10")
 
    # =====================================================
    # DCF
    # =====================================================
    wd = wb.create_sheet("DCF")
    wd["A1"] = "DCF企業価値評価 (百万円)"
    wd["A1"].font = TITLE
    _header(wd, 2, YEARS[1:])
    _label(wd, 3, "期間 (年)")
    _label(wd, 4, "フリーキャッシュフロー (FCF)")
    _label(wd, 5, "割引係数")
    _label(wd, 6, "FCFの現在価値")
    for i in range(NUM_YEARS):
        col = get_column_letter(2 + i)
        cf_col = get_column_letter(3 + i)
        _put(wd, f"{col}3", i + 1, BLUE, "0")
        _put(wd, f"{col}4", f"={C}!{cf_col}14", GREEN)
        _put(wd, f"{col}5", f"=(1+{A}!$B$12)^{col}3", BLACK, "0.0000")
        _put(wd, f"{col}6", f"={col}4/{col}5")
 
    labels = {
        8: "【永続成長法】", 9: "FCFの現在価値合計", 10: "ターミナルバリュー",
        11: "ターミナルバリューの現在価値", 12: "企業価値 (EV)",
        13: "純有利子負債 (Yr0)", 14: "株主価値", 15: "理論株価 (円)",
        17: "【エグジットマルチプル法】", 18: "ターミナルバリュー (EBITDA×倍率)",
        19: "ターミナルバリューの現在価値", 20: "企業価値 (EV)",
        21: "株主価値", 22: "理論株価 (円)",
    }
    for r, lab in labels.items():
        _label(wd, r, lab, bold=r in (8, 12, 15, 17, 20, 22))
    _put(wd, "B9", "=SUM(B6:F6)")
    _put(wd, "B10", f"=F4*(1+{A}!B13)/({A}!B12-{A}!B13)")
    _put(wd, "B11", "=B10/(1+" + f"{A}!B12)^5")
    _put(wd, "B12", "=B9+B11")
    _put(wd, "B13", f"={B}!B9-{B}!B3", GREEN)
    _put(wd, "B14", "=B12-B13")
    _put(wd, "B15", f"=B14*100/{A}!B15", BLACK, '#,##0')
    _put(wd, "B18", f"=({I}!G7+{C}!G4)*{A}!B14")
    _put(wd, "B19", f"=B18/(1+{A}!B12)^5")
    _put(wd, "B20", "=B9+B19")
    _put(wd, "B21", "=B20-B13")
    _put(wd, "B22", f"=B21*100/{A}!B15", BLACK, '#,##0')
 
    for sheet in (wi, wbs, wc, wd):
        sheet.column_dimensions["A"].width = 38
        for i in range(2, 8):
            sheet.column_dimensions[get_column_letter(i)].width = 14
        sheet.freeze_panes = "B3"
 
    # =====================================================
    # 値のまま貼り付けるシート (レーシオ・感度分析など)
    # =====================================================
    for name, df in (extra_sheets or {}).items():
        wx = wb.create_sheet(name[:31])
        wx.cell(row=1, column=1, value=name).font = TITLE
        for j, col_name in enumerate(df.columns):
            c = wx.cell(row=3, column=2 + j, value=str(col_name))
            c.font = HEAD_FONT
            c.fill = HEAD_FILL
        for i, (idx, row) in enumerate(df.iterrows()):
            wx.cell(row=4 + i, column=1, value=str(idx)).font = BOLD
            for j, v in enumerate(row):
                if pd.isna(v):
                    v = None
                cell = wx.cell(row=4 + i, column=2 + j, value=v)
                cell.font = BLACK
                cell.number_format = "#,##0.00"
        wx.column_dimensions["A"].width = 30
        for j in range(len(df.columns)):
            wx.column_dimensions[get_column_letter(2 + j)].width = 16
 
    return wb
 
 
def build_excel_bytes(p, extra_sheets=None):
    wb = build_workbook(p, extra_sheets)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
 